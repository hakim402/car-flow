import csv
import hashlib
import io
import json
import os
from pathlib import Path
from datetime import date, timedelta
from decimal import Decimal
from unittest.mock import patch

import pytest
from django.urls import reverse
from django.utils import timezone, translation

from apps.accounts.models import Role
from apps.accounting.exports import (
    export_csv, export_docx, export_pdf, export_xlsx, freeze_report,
    pdf_html, public_snapshot, safe_text,
)
from apps.accounting.models import ReportExport, ReportShare, ReportShareAccess
from apps.accounting.tasks import cleanup_report_delivery, generate_report_export
from apps.core.testing import OrganizationFactory, UserFactory


@pytest.fixture
def report_data():
    return {
        "key": "activity", "title": "Business Activity", "description": "Report description",
        "columns": [{"key": "name", "label": "Customer", "kind": "text"},
                    {"key": "amount", "label": "Outstanding", "kind": "money"},
                    {"key": "currency", "label": "Currency", "kind": "text"},
                    {"key": "date", "label": "Due date", "kind": "date"},
                    {"key": "phone", "label": "Phone", "kind": "text"}],
        "rows": [{"name": f"Customer {i}", "amount": Decimal("12.34"), "currency": "USD",
                  "date": date(2026, 9, 1), "phone": "+93700000000", "_url": "/customers/1/",
                  "_image": "http://internal/private.jpg"} for i in range(30)],
        "metrics": [{"label": "Outstanding", "value": Decimal("370.20"), "currency": "USD"}],
        "charts": [{"title": "Balances", "currency": "USD", "items": [{"label": "Private label", "value": 12, "url": "/customers/1/"}]}],
        "notes": ["Internal note"],
    }


@pytest.fixture
def snapshot(report_data):
    return freeze_report(report_data, company_name="Example company", language="en")


@pytest.fixture
def manager(db):
    user = UserFactory()
    user.roles.add(Role.objects.get(key="org_admin"))
    return user


def test_freeze_preserves_exact_amounts_and_strips_internal_navigation(snapshot):
    assert snapshot["rows"][0]["amount"] == "12.34"
    assert snapshot["rows"][0]["date"] == "2026-09-01"
    assert "_url" not in snapshot["rows"][0]
    assert "_image" not in snapshot["rows"][0]
    assert "url" not in snapshot["charts"][0]["items"][0]


def test_share_only_selected_columns_and_automatic_currency(snapshot):
    result = public_snapshot(snapshot, ["amount"])
    assert {c["key"] for c in result["columns"]} == {"amount", "currency"}
    assert set(result["rows"][0]) == {"amount", "currency"}
    assert result["charts"] == result["metrics"] == result["notes"] == []
    assert result["filters"] == {}


@pytest.mark.parametrize("columns", [[], ["phone"], ["password"], ["_url"], ["nonexistent"]])
def test_share_rejects_unapproved_or_empty_columns(snapshot, columns):
    with pytest.raises(ValueError):
        public_snapshot(snapshot, columns)


@pytest.mark.parametrize("value", ["=1+1", "+CMD()", "-1+1", "@SUM(1)", " \t=HYPERLINK('x')"])
def test_spreadsheet_formula_injection_is_escaped(value):
    assert safe_text(value).startswith("'")


def test_csv_includes_all_rows_with_utf8_and_escaped_cells(snapshot):
    snapshot["rows"][0]["name"] = "=HYPERLINK(\"https://example.com\")"
    snapshot["rows"][1]["name"] = "د پېرودونکو حسابونه"
    data = export_csv(snapshot)
    assert data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert len(rows) == 31
    assert rows[1][0].startswith("'=")
    assert rows[2][0] == "د پېرودونکو حسابونه"


def test_xlsx_typed_data_rtl_and_formula_safety(snapshot):
    from openpyxl import load_workbook
    snapshot["language"] = "ps"
    snapshot["rows"][0]["name"] = "=1+1"
    with translation.override("en"):
        workbook = load_workbook(io.BytesIO(export_xlsx(snapshot)))
    sheet = workbook["Records"]
    assert sheet.max_row == 31
    assert sheet["A2"].data_type == "s"
    assert sheet["B2"].value == 12.34
    assert sheet["D2"].value.date() == date(2026, 9, 1)
    assert sheet.sheet_view.rightToLeft
    assert sheet.freeze_panes == "A2"


def test_docx_contains_all_rows_and_rtl_semantics(snapshot):
    from docx import Document
    snapshot["language"] = "prs"
    snapshot["title"] = "گزارش حساب مشتریان"
    document = Document(io.BytesIO(export_docx(snapshot)))
    assert len(document.tables[0].rows) == 31
    assert "w:bidiVisual" in document.tables[0]._tbl.xml
    assert "w:bidi" in document.paragraphs[1]._p.xml
    assert any(snapshot["title"] in p.text for p in document.paragraphs)


@pytest.mark.parametrize("language,title", [("en", "Cash Position"), ("prs", "موقف نقدی"), ("ps", "د نغدو موقف")])
def test_pdf_renders_three_languages_without_network(snapshot, language, title):
    snapshot["language"] = language
    snapshot["title"] = title
    snapshot["rows"][0]["name"] = '<img src="http://127.0.0.1:9/private">'
    markup = pdf_html(snapshot)
    assert "&lt;img" in markup
    assert "<img" not in markup
    assert "dir='rtl'" in markup if language != "en" else "dir='ltr'" in markup
    data = export_pdf(snapshot)
    assert data.startswith(b"%PDF-")
    if os.environ.get("REPORT_QA_OUTPUT_DIR"):
        output = Path(os.environ["REPORT_QA_OUTPUT_DIR"])
        output.mkdir(parents=True, exist_ok=True)
        (output / f"report-{language}.pdf").write_bytes(data)


@pytest.mark.django_db
def test_export_requires_explicit_permission_even_for_unassigned_user(client):
    client.force_login(UserFactory())
    assert client.post(reverse("accounting:export_create", args=["activity"]), {"format": "csv"}).status_code == 403


@pytest.mark.django_db
def test_export_snapshot_download_and_owner_isolation(client, manager, report_data):
    client.force_login(manager)
    with patch("apps.accounting.export_views.build_report", return_value=report_data):
        response = client.post(reverse("accounting:export_create", args=["activity"]), {"format": "csv"})
    assert response.status_code == 302
    job = ReportExport.all_objects.get(created_by=manager)
    assert job.status == "ready"
    assert job.row_count == 30
    assert job.file_content
    response = client.get(reverse("accounting:export_download", args=[job.pk]))
    assert response.status_code == 200
    assert "no-store" in response["Cache-Control"]
    assert len(list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))) == 31
    other = UserFactory(company=manager.company)
    other.roles.add(Role.objects.get(key="org_admin"))
    client.force_login(other)
    assert client.get(reverse("accounting:export_download", args=[job.pk])).status_code == 404
    client.force_login(UserFactory())
    assert client.get(reverse("accounting:export_status", args=[job.pk])).status_code == 404


@pytest.mark.django_db
def test_export_status_uses_catalog_title_when_snapshot_title_missing(client, manager):
    job = ReportExport.all_objects.create(company=manager.company, created_by=manager,
        report_key="activity", format="pdf", snapshot={}, row_count=0, status="pending",
        expires_at=timezone.now() + timedelta(hours=1))
    client.force_login(manager)
    response = client.get(reverse("accounting:export_status", args=[job.pk]))
    assert response.status_code == 200
    assert "Business Activity" in response.content.decode()


@pytest.mark.django_db
def test_export_rechecks_permission_and_expiration(client, manager, snapshot):
    job = ReportExport.all_objects.create(company=manager.company, created_by=manager,
        report_key="activity", format="csv", snapshot=snapshot, file_content=b"private", status="ready",
        expires_at=timezone.now() + timedelta(hours=1))
    client.force_login(manager)
    manager.roles.clear()
    assert client.get(reverse("accounting:export_download", args=[job.pk])).status_code == 403
    manager.roles.add(Role.objects.get(key="org_admin"))
    ReportExport.all_objects.filter(pk=job.pk).update(expires_at=timezone.now() - timedelta(seconds=1))
    assert client.get(reverse("accounting:export_download", args=[job.pk])).status_code == 410
    job.refresh_from_db()
    assert not job.file_content and not job.snapshot


@pytest.mark.django_db
def test_share_token_is_hashed_snapshot_only_and_revocable(client, manager, report_data):
    client.force_login(manager)
    with patch("apps.accounting.export_views.build_report", return_value=report_data):
        response = client.post(reverse("accounting:share_create", args=["activity"]),
                               {"columns": ["name", "amount"], "expires_in": "3600"})
    assert response.status_code == 201
    url = response.context["share_url"]
    token = url.rstrip("/").split("/")[-1]
    share = ReportShare.all_objects.get(created_by=manager)
    assert share.token_hash == hashlib.sha256(token.encode()).hexdigest()
    assert token not in json.dumps(share.snapshot)
    assert "phone" not in share.snapshot["rows"][0]
    assert share.snapshot["metrics"] == []
    client.logout()
    response = client.get(url)
    assert response.status_code == 200
    assert "no-store" in response["Cache-Control"]
    assert response["Referrer-Policy"] == "no-referrer"
    assert "script-src 'none'" in response["Content-Security-Policy"]
    assert len(response.context["display_rows"]) == 25
    assert ReportShareAccess.all_objects.filter(share=share).count() == 1
    client.force_login(manager)
    response = client.post(reverse("accounting:share_revoke", args=[share.pk]))
    assert response.status_code == 302
    client.logout()
    assert client.get(url).status_code == 404


@pytest.mark.django_db
def test_share_invalid_when_creator_loses_access_or_tenant_changes(client, manager, snapshot):
    token = "a" * 43
    share = ReportShare.all_objects.create(company=manager.company, created_by=manager,
        report_key="activity", token_hash=hashlib.sha256(token.encode()).hexdigest(),
        snapshot=public_snapshot(snapshot, ["name"]), expires_at=timezone.now() + timedelta(hours=1))
    url = reverse("accounting:shared_report", args=[token])
    assert client.get(url).status_code == 200
    manager.is_active = False; manager.save()
    assert client.get(url).status_code == 404
    manager.is_active = True; manager.company = OrganizationFactory(); manager.save()
    assert client.get(url).status_code == 404


@pytest.mark.django_db
def test_delivery_filters_and_post_only_endpoints(client, manager):
    client.force_login(manager)
    url = reverse("accounting:export_create", args=["activity"])
    assert client.get(url).status_code == 405
    assert client.post(url + "?date_from=not-a-date", {"format": "csv"}).status_code == 400
    assert client.post(url + "?branch=99999999", {"format": "csv"}).status_code == 400
    assert client.post(url, {"format": "exe"}).status_code == 400
    assert not ReportExport.all_objects.exists()


@pytest.mark.django_db
def test_cleanup_removes_payloads_not_business_records(manager, snapshot):
    expired = timezone.now() - timedelta(seconds=1)
    export = ReportExport.all_objects.create(company=manager.company, created_by=manager,
        report_key="activity", format="csv", snapshot=snapshot, status="ready",
        file_content=b"private", expires_at=expired)
    share = ReportShare.all_objects.create(company=manager.company, created_by=manager,
        report_key="activity", token_hash="a" * 64, snapshot=snapshot, expires_at=expired)
    cleanup_report_delivery()
    export.refresh_from_db(); share.refresh_from_db()
    assert export.status == "expired" and not export.snapshot and not export.file_content
    assert not share.snapshot


@pytest.mark.django_db
def test_background_job_rejects_changed_creator_company(manager, snapshot):
    job = ReportExport.all_objects.create(company=manager.company, created_by=manager,
        report_key="activity", format="csv", snapshot=snapshot,
        expires_at=timezone.now() + timedelta(hours=1))
    manager.company = OrganizationFactory(); manager.save()
    generate_report_export(job.pk)
    job.refresh_from_db()
    assert job.status == "failed"
    assert not job.file_content
