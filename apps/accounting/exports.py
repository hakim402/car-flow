"""Serialize one immutable report dataset into safe, multilingual files."""
import csv
import html
import io
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.core.serializers.json import DjangoJSONEncoder
from django.utils import timezone
from django.utils.translation import gettext as _

FORMATS = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
MAX_SNAPSHOT_BYTES = 15 * 1024 * 1024
MAX_FILE_BYTES = 25 * 1024 * 1024


def filter_label(key):
    return {
        "date_from": _("From date"), "date_to": _("To date"),
        "as_of": _("As of date"), "branch": _("Branch"),
        "currency": _("Currency"), "q": _("Search"), "category": _("Category"),
    }.get(key, key)


def freeze_report(report, *, company_name, language, filters=None):
    """Remove navigation/media metadata, freeze lazy translations and decimals."""
    def clean(value):
        if isinstance(value, dict):
            return {str(k): clean(v) for k, v in value.items()
                    if not str(k).startswith("_") and k not in {"url", "image"}}
        if isinstance(value, (list, tuple)):
            return [clean(v) for v in value]
        return value
    frozen = clean(report)
    frozen.update(company_name=str(company_name), language=language,
                  generated_at=timezone.now().isoformat(), filters=filters or {})
    encoded = json.dumps(frozen, cls=DjangoJSONEncoder, ensure_ascii=False)
    if len(encoded.encode("utf-8")) > getattr(settings, "REPORT_MAX_SNAPSHOT_BYTES", MAX_SNAPSHOT_BYTES):
        raise ValueError(_("This report is too large. Narrow the filters and try again."))
    return json.loads(encoded)


def shareable_columns(report):
    """Positive report column schema only; never disclose sensitive text fields."""
    private = {"phone", "email", "address", "notes", "body", "raw_payload", "identity_number",
               "tazkira", "passport", "credentials", "token", "password"}
    return [c for c in report["columns"] if c.get("shareable", True)
            and c["key"] not in private and not c["key"].startswith("_")]


def public_snapshot(snapshot, selected):
    allowed = {c["key"] for c in shareable_columns(snapshot)}
    if not selected or not set(selected).issubset(allowed):
        raise ValueError(_("Select at least one approved report column."))
    selected = set(selected)
    if "currency" in allowed and any(c["key"] in selected and c.get("kind") == "money" for c in snapshot["columns"]):
        selected.add("currency")
    # Do not leak hidden-column values indirectly via metrics, chart labels,
    # search terms, internal warnings, links, or notes.
    return {
        "key": snapshot["key"], "title": snapshot["title"],
        "description": snapshot.get("description", ""),
        "company_name": snapshot["company_name"], "language": snapshot["language"],
        "generated_at": snapshot["generated_at"],
        "columns": [c for c in snapshot["columns"] if c["key"] in selected],
        "rows": [{c["key"]: row.get(c["key"], "") for c in snapshot["columns"]
                  if c["key"] in selected} for row in snapshot["rows"]],
        "metrics": [], "charts": [], "notes": [], "filters": {},
    }


def safe_text(value):
    """Prevent formula evaluation when user data is opened by spreadsheet apps."""
    text = "" if value is None else str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r", "\n")) else text


def display(value, kind="text"):
    if value is None:
        return "—"
    if kind == "money":
        try:
            return f"{Decimal(str(value)):,.2f}"
        except InvalidOperation:
            pass
    return str(value)


def export_csv(snapshot):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([safe_text(c["label"]) for c in snapshot["columns"]])
    for row in snapshot["rows"]:
        writer.writerow([safe_text(row.get(c["key"])) for c in snapshot["columns"]])
    # UTF-8 BOM keeps Dari/Pashto text readable in Excel's CSV opener.
    return output.getvalue().encode("utf-8-sig")


def export_xlsx(snapshot):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    workbook = Workbook()
    summary = workbook.active
    summary.title = str(_("Summary"))[:31]
    summary.append(["AMOXRUNS", safe_text(snapshot["title"])])
    summary.append([str(_("Company")), safe_text(snapshot["company_name"])])
    summary.append([str(_("Generated at")), snapshot["generated_at"]])
    for key, value in snapshot.get("filters", {}).items():
        if value:
            summary.append([safe_text(filter_label(key)), safe_text(value)])
    summary.append([])
    summary.append([str(_("Metric")), str(_("Value")), str(_("Currency"))])
    for metric in snapshot.get("metrics", []):
        try:
            value = Decimal(str(metric["value"]))
        except InvalidOperation:
            value = safe_text(metric["value"])
        summary.append([safe_text(metric["label"]), value, safe_text(metric.get("currency", ""))])
    for note in snapshot.get("notes", []):
        summary.append([safe_text(note)])
    sheet = workbook.create_sheet(str(_("Records"))[:31])
    sheet.append([safe_text(c["label"]) for c in snapshot["columns"]])
    for row in snapshot["rows"]:
        values = []
        for col in snapshot["columns"]:
            value = row.get(col["key"])
            if value not in (None, "") and col.get("kind") in {"money", "number"}:
                try:
                    value = Decimal(str(value))
                except InvalidOperation:
                    value = safe_text(value)
            elif value and col.get("kind") == "date":
                try:
                    value = datetime.fromisoformat(str(value)).replace(tzinfo=None)
                except ValueError:
                    value = safe_text(value)
            elif value is not None:
                value = safe_text(value)
            values.append(value)
        sheet.append(values)
    for column, spec in enumerate(snapshot["columns"], 1):
        if spec.get("kind") in {"money", "date"}:
            for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
                cells[0].number_format = "#,##0.00" if spec["kind"] == "money" else "yyyy-mm-dd"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    for index, chart in enumerate(snapshot.get("charts", []), 1):
        chart_sheet = workbook.create_sheet(f"{_('Chart')} {index}"[:31])
        chart_sheet.append([safe_text(chart["title"]), safe_text(chart.get("currency", ""))])
        for item in chart.get("items", []):
            chart_sheet.append([safe_text(item["label"]), Decimal(str(item["value"]))])
        if chart_sheet.max_row > 1:
            bars = BarChart()
            bars.type = "bar"
            bars.title = chart["title"]
            bars.add_data(Reference(chart_sheet, min_col=2, min_row=1, max_row=chart_sheet.max_row), titles_from_data=True)
            bars.set_categories(Reference(chart_sheet, min_col=1, min_row=2, max_row=chart_sheet.max_row))
            chart_sheet.add_chart(bars, "D2")
    for ws in workbook:
        ws.sheet_view.rightToLeft = snapshot.get("language") in {"prs", "ps"}
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
            cell.font = Font(name="Noto Sans Arabic", color="FFFFFF", bold=True)
        for row in ws:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row > 1:
                    cell.font = Font(name="Noto Sans Arabic", size=11)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 24
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def export_docx(snapshot):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    document = Document()
    section = document.sections[0]
    if len(snapshot["columns"]) > 5:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Inches(.5)
    document.styles["Normal"].font.name = "Noto Sans Arabic"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading("AMOXRUNS", 0)
    document.add_heading(snapshot["title"], 1)
    document.add_paragraph(snapshot["company_name"])
    document.add_paragraph(f"{_('Generated at')}: {snapshot['generated_at']}")
    for key, value in snapshot.get("filters", {}).items():
        if value:
            document.add_paragraph(f"{filter_label(key)}: {value}")
    for metric in snapshot.get("metrics", []):
        document.add_paragraph(f"{metric['label']}: {display(metric['value'], 'money' if metric.get('currency') else 'number')} {metric.get('currency', '')}")
    for note in snapshot.get("notes", []):
        document.add_paragraph(str(note))
    table = document.add_table(rows=1, cols=len(snapshot["columns"])); table.style = "Light Shading Accent 1"
    for cell, col in zip(table.rows[0].cells, snapshot["columns"]):
        cell.text = str(col["label"])
    header_repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header_repeat)
    for row in snapshot["rows"]:
        for cell, col in zip(table.add_row().cells, snapshot["columns"]):
            cell.text = display(row.get(col["key"]), col.get("kind"))
    rtl = snapshot.get("language") in {"prs", "ps"}
    if rtl:
        bidi_table = OxmlElement("w:bidiVisual"); table._tbl.tblPr.append(bidi_table)
    paragraphs = list(document.paragraphs) + [p for row in table.rows for cell in row.cells for p in cell.paragraphs]
    for paragraph in paragraphs:
        if rtl:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
        for run in paragraph.runs:
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:cs"), "Noto Sans Arabic")
            if rtl:
                run._element.get_or_add_rPr().append(OxmlElement("w:rtl"))
    stream = io.BytesIO(); document.save(stream)
    return stream.getvalue()


def pdf_html(snapshot):
    """Only escaped text and server-owned CSS: no remote images, links or JS."""
    esc = lambda value: html.escape(str(value), quote=True)
    rtl = snapshot.get("language") in {"prs", "ps"}
    header = "".join(f"<th>{esc(c['label'])}</th>" for c in snapshot["columns"])
    records = "".join("<tr>" + "".join(
        f"<td>{esc(display(row.get(c['key']), c.get('kind')))}</td>"
        for c in snapshot["columns"]) + "</tr>" for row in snapshot["rows"])
    metrics = "".join(f"<div class='metric'><span>{esc(m['label'])}</span><strong>{esc(display(m['value'], 'money' if m.get('currency') else 'number'))} {esc(m.get('currency', ''))}</strong></div>" for m in snapshot.get("metrics", []))
    notes = "".join(f"<p>{esc(n)}</p>" for n in snapshot.get("notes", []))
    filters = " · ".join(f"{esc(filter_label(k))}: {esc(v)}" for k, v in snapshot.get("filters", {}).items() if v)
    charts = []
    for chart in snapshot.get("charts", []):
        items = chart.get("items", [])
        maximum = max([abs(Decimal(str(i["value"]))) for i in items] or [Decimal(1)]) or Decimal(1)
        bars = "".join(f"<div class='chart-row'><span>{esc(i['label'])}</span><div class='bar' style='width:{min(100, abs(Decimal(str(i['value']))) / maximum * 100):.2f}%'></div><b>{esc(display(i['value'], 'money'))}</b></div>" for i in items)
        charts.append(f"<section><h2>{esc(chart['title'])} {esc(chart.get('currency', ''))}</h2>{bars}</section>")
    return f"""<!doctype html><html lang='{esc(snapshot.get('language', 'en'))}' dir='{'rtl' if rtl else 'ltr'}'>
    <head><meta charset='utf-8'><meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; style-src 'unsafe-inline'\"><title>{esc(snapshot['title'])}</title><style>
    @page {{ size: A4 landscape; margin: 12mm; }}
    body {{font:10px 'Noto Sans Arabic','Noto Sans',sans-serif;color:#17233b;line-height:1.6}}
    h1 {{font-size:23px;margin:6px 0}} h2 {{font-size:13px;color:#1d4ed8}} .brand {{color:#1d4ed8;font-weight:bold;font-size:17px}}
    .meta {{color:#526079}} .metrics {{display:flex;gap:10px;flex-wrap:wrap;margin:15px 0}}
    .metric {{border:1px solid #dbe3ef;border-radius:6px;padding:8px;min-width:100px}} .metric span,.metric strong {{display:block}}
    table {{width:100%;border-collapse:collapse;table-layout:fixed;margin-top:15px}} thead {{display:table-header-group}} tr {{break-inside:avoid}}
    th,td {{text-align:{'right' if rtl else 'left'};padding:7px;border-bottom:1px solid #dbe3ef;overflow-wrap:anywhere;vertical-align:top}} th {{background:#1d4ed8;color:white}}
    tbody tr:nth-child(even) {{background:#f3f6fc}} section {{break-inside:avoid;margin:15px 0}} .chart-row {{display:grid;grid-template-columns:150px 1fr 100px;gap:8px;align-items:center;margin:5px 0}} .bar {{height:10px;background:#2563eb;border-radius:3px}}
    </style></head><body><div class='brand'>AMOXRUNS</div><h1>{esc(snapshot['title'])}</h1><div class='meta'>{esc(snapshot['company_name'])} · {esc(_('Generated at'))}: {esc(snapshot['generated_at'])}</div><p class='meta'>{filters}</p><div class='metrics'>{metrics}</div>{notes}{''.join(charts)}<table><thead><tr>{header}</tr></thead><tbody>{records}</tbody></table></body></html>"""


def export_pdf(snapshot):
    # Browser-quality bidi and shaping (unlike renderers lacking RTL support).
    # This browser never opens application URLs and has no authenticated state.
    from playwright.sync_api import sync_playwright
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(
            executable_path=getattr(settings, "REPORT_CHROMIUM_PATH", "/usr/bin/chromium-headless-shell"),
            headless=True, args=["--disable-dev-shm-usage"],
        )
        try:
            context = browser.new_context(java_script_enabled=False, offline=True, service_workers="block")
            context.route("**/*", lambda route: route.abort())
            page = context.new_page()
            page.set_content(pdf_html(snapshot), wait_until="load", timeout=30000)
            return page.pdf(format="A4", landscape=True, print_background=True, prefer_css_page_size=True)
        finally:
            browser.close()


def render_export(snapshot, format):
    renderer = {"csv": export_csv, "xlsx": export_xlsx, "docx": export_docx, "pdf": export_pdf}[format]
    data = renderer(snapshot)
    if len(data) > getattr(settings, "REPORT_MAX_FILE_BYTES", MAX_FILE_BYTES):
        raise ValueError(_("This report is too large. Narrow the filters and try again."))
    return data
